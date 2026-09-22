from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.db.models import Sum, F, ExpressionWrapper, DecimalField
from django.db.models.functions import Coalesce
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from accounts.permissions import IsOwner
from parametres.models import ParametresBoutique
from tenants.mixins import BoutiqueScopedMixin
from sales.models import LigneVente, Vente
from sales.utils import unites_reelles_expr
from purchases.models import Achat
from expenses.models import Depense

# Palette du PDF - reprend l'identité visuelle déjà utilisée dans
# Reports.jsx/Dashboard.jsx (bleu Tailwind par défaut), pas une palette
# inventée pour l'occasion.
PDF_BLEU_PRINCIPAL = colors.HexColor('#2563eb')   # blue-600
PDF_GRIS_TITRE = colors.HexColor('#0f172a')       # slate-900
PDF_GRIS_TEXTE = colors.HexColor('#64748b')       # slate-500
PDF_GRIS_LIGNE_ALTERNEE = colors.HexColor('#f1f5f9')  # slate-100
PDF_GRIS_BORDURE = colors.HexColor('#e2e8f0')     # slate-200
PDF_VERT_POSITIF = colors.HexColor('#059669')     # emerald-600
PDF_ROUGE_NEGATIF = colors.HexColor('#dc2626')    # red-600


def calculer_resume_financier(boutique, date_debut, date_fin):
    """Calcule le résumé financier (CA, achats, dépenses, bénéfices) d'une
    boutique sur une période optionnelle. Extrait de ResumeFinancierView.get()
    pour être réutilisé par ResumeFinancierExportPDFView - un seul endroit à
    corriger si la logique de calcul change (même principe que le bénéfice
    historique figé, cf. sales.models.LigneVente.prix_achat_unitaire)."""
    # Les ventes annulées ne doivent plus compter dans le CA ni le
    # bénéfice (cf. VenteViewSet.annuler) ; benefice_brut hérite de ce
    # filtre via `vente__in=ventes_qs` plus bas.
    ventes_qs = Vente.objects.filter(boutique=boutique, statut='VALIDEE')
    # Les achats annulés ne doivent plus compter dans le total ni le
    # nombre d'achats (cf. AchatViewSet.annuler).
    achats_qs = Achat.objects.filter(boutique=boutique, statut='VALIDE')
    # Les dépenses annulées ne doivent plus compter dans le total ni le
    # bénéfice net (cf. DepenseViewSet.annuler, P2 point 15).
    depenses_qs = Depense.objects.filter(boutique=boutique, statut='VALIDEE')

    if date_debut and date_fin:
        ventes_qs = ventes_qs.filter(date_vente__date__range=[date_debut, date_fin])
        achats_qs = achats_qs.filter(date_achat__date__range=[date_debut, date_fin])
        depenses_qs = depenses_qs.filter(date_depense__range=[date_debut, date_fin])

    # Chiffre d'affaires total
    chiffre_affaires = ventes_qs.aggregate(total=Sum('montant_net'))['total'] or 0

    # Total des achats
    total_achats = achats_qs.aggregate(total=Sum('montant_total'))['total'] or 0

    # Total des dépenses
    total_depenses = depenses_qs.aggregate(total=Sum('montant'))['total'] or 0

    # --- Bénéfice brut, calculé directement en base de données ---
    # Convertit chaque ligne de vente en nombre d'unités réelles via le
    # facteur de conversion centralisé sur son unité de vente.
    unites_reelles = unites_reelles_expr()

    # Le coût d'achat HISTORIQUE (prix_achat_unitaire x unités réelles)
    # est figé au moment de la vente - jamais Produit.prix_achat courant,
    # qui change à chaque nouvel achat et déformerait rétroactivement le
    # bénéfice d'une vente déjà réalisée (bug P1 corrigé). Le repli sur
    # Produit.prix_achat via Coalesce ne s'applique qu'aux lignes créées
    # avant ce correctif (prix_achat_unitaire NULL, aucun coût
    # historique connu) - limite assumée et documentée, voir
    # sales.models.LigneVente.prix_achat_unitaire.
    cout_historique_expr = ExpressionWrapper(
        Coalesce(F('prix_achat_unitaire'), F('produit__prix_achat')) * unites_reelles,
        output_field=DecimalField(max_digits=14, decimal_places=2)
    )

    cout_historique_total = LigneVente.objects.filter(
        boutique=boutique,
        vente__in=ventes_qs
    ).aggregate(total=Sum(cout_historique_expr))['total'] or 0

    # Le CA est la somme des montants nets : une remise globale diminue
    # donc le bénéfice brut au même titre que le chiffre d'affaires.
    benefice_brut = chiffre_affaires - cout_historique_total

    # Bénéfice net = Bénéfice brut - Total des dépenses
    benefice_net = benefice_brut - total_depenses

    return {
        "date_debut": date_debut,
        "date_fin": date_fin,
        "chiffre_affaires": chiffre_affaires,
        "total_achats": total_achats,
        "total_depenses": total_depenses,
        "benefice_brut": benefice_brut,
        "benefice_net": benefice_net,
        "nombre_ventes": ventes_qs.count(),
        "nombre_achats": achats_qs.count(),
        "nombre_depenses": depenses_qs.count(),
    }


def parametres_boutique(boutique):
    """get_or_create comme ParametresBoutiqueView.get_object() : une
    boutique n'a pas forcément encore de ParametresBoutique créé
    explicitement (default="FCFA" s'applique alors). Réutilisé par les
    vues d'export (PDF, Excel) - une seule source pour ce pattern."""
    parametres, _ = ParametresBoutique.objects.get_or_create(boutique=boutique)
    return parametres


class ResumeFinancierView(BoutiqueScopedMixin, APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, *args, **kwargs):
        # Réutilise la même résolution de boutique effective (Vue Support
        # comprise) que les ViewSets scopés, au lieu de dupliquer cette
        # logique ici sans gestion propre de Boutique.DoesNotExist (P2
        # point 14 - Vue Support ad-hoc). Vue en lecture seule : pas
        # d'appel à _verifier_acces() ici - consulter ses rapports doit
        # rester possible boutique désactivée/abonnement expiré, seules
        # les écritures sont bloquées (cf. tenants.mixins.BoutiqueScopedMixin).
        boutique = self._boutique_effective()
        # Récupérer les filtres de date optionnels (?date_debut=YYYY-MM-DD&date_fin=YYYY-MM-DD)
        date_debut = request.GET.get('date_debut')
        date_fin = request.GET.get('date_fin')

        return Response(calculer_resume_financier(boutique, date_debut, date_fin))


class ResumeFinancierExportPDFView(BoutiqueScopedMixin, APIView):
    # Export réservé au propriétaire (même principe que VenteViewSet.annuler
    # - une écriture comptable... ici une extraction de données comptables
    # complètes, pas un simple affichage - P1 point 6, RBAC).
    permission_classes = [IsAuthenticated, IsOwner]

    def get(self, request, *args, **kwargs):
        boutique = self._boutique_effective()
        date_debut = request.GET.get('date_debut')
        date_fin = request.GET.get('date_fin')
        resultat = calculer_resume_financier(boutique, date_debut, date_fin)
        parametres = parametres_boutique(boutique)
        devise = parametres.devise

        response = HttpResponse(content_type='application/pdf')
        nom_fichier = f"resume-financier-{boutique.slug}-{timezone.localdate().isoformat()}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'

        marge = 2 * cm
        # pageCompression=0 : flux non compressé - permet aux tests de
        # vérifier le contenu textuel directement dans response.content
        # sans dépendance supplémentaire d'extraction PDF (pypdf, etc.).
        doc = SimpleDocTemplate(
            response, pagesize=A4, pageCompression=0,
            leftMargin=marge, rightMargin=marge, topMargin=marge, bottomMargin=marge,
        )
        largeur_disponible = A4[0] - 2 * marge

        elements = [
            self._entete(boutique, parametres, resultat, largeur_disponible),
            Spacer(1, 0.8 * cm),
            self._tableau_resume(resultat, devise, largeur_disponible),
        ]
        doc.build(elements, onFirstPage=self._pied_de_page, onLaterPages=self._pied_de_page)
        return response

    def _entete(self, boutique, parametres, resultat, largeur_disponible):
        styles = getSampleStyleSheet()
        style_titre = ParagraphStyle(
            'TitreBoutique', parent=styles['Normal'],
            fontName='Helvetica-Bold', fontSize=18, textColor=PDF_GRIS_TITRE,
        )
        style_periode = ParagraphStyle(
            'Periode', parent=styles['Normal'],
            fontName='Helvetica', fontSize=10, textColor=PDF_GRIS_TEXTE, spaceBefore=8,
        )

        if resultat['date_debut'] and resultat['date_fin']:
            texte_periode = f"Période : du {resultat['date_debut']} au {resultat['date_fin']}"
        else:
            texte_periode = "Période : toutes dates confondues"

        bloc_titre = [
            Paragraph(f"Résumé financier — {boutique.nom}", style_titre),
            Paragraph(texte_periode, style_periode),
        ]

        logo_flowable = self._logo_flowable(parametres)
        if logo_flowable is not None:
            largeur_logo_colonne = 3 * cm
            entete = Table(
                [[logo_flowable, bloc_titre]],
                colWidths=[largeur_logo_colonne, largeur_disponible - largeur_logo_colonne],
            )
        else:
            entete = Table([[bloc_titre]], colWidths=[largeur_disponible])

        entete.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        return entete

    def _logo_flowable(self, parametres):
        # ParametresBoutique.logo est optionnel - et même s'il est
        # renseigné en base, le fichier peut être absent du stockage
        # (même limite documentée pour l'affichage du logo côté frontend,
        # cf. Settings.jsx/logoIndisponible) : jamais d'erreur, juste pas
        # de logo dans le PDF le cas échéant.
        if not parametres.logo:
            return None
        try:
            with parametres.logo.open('rb') as fichier_logo:
                lecteur_image = ImageReader(BytesIO(fichier_logo.read()))
        except (OSError, ValueError):
            return None

        largeur_logo, hauteur_logo = lecteur_image.getSize()
        if not largeur_logo or not hauteur_logo:
            return None

        hauteur_cible = 1.5 * cm
        largeur_cible = largeur_logo * (hauteur_cible / hauteur_logo)
        return Image(lecteur_image, width=largeur_cible, height=hauteur_cible)

    def _tableau_resume(self, resultat, devise, largeur_disponible):
        lignes = [
            ("Chiffre d'affaires", f"{resultat['chiffre_affaires']:.2f} {devise}"),
            ("Total des achats", f"{resultat['total_achats']:.2f} {devise}"),
            ("Total des dépenses", f"{resultat['total_depenses']:.2f} {devise}"),
            ("Bénéfice brut", f"{resultat['benefice_brut']:.2f} {devise}"),
            ("Bénéfice net", f"{resultat['benefice_net']:.2f} {devise}"),
            ("Nombre de ventes", str(resultat['nombre_ventes'])),
            ("Nombre d'achats", str(resultat['nombre_achats'])),
            ("Nombre de dépenses", str(resultat['nombre_depenses'])),
        ]
        # Index (dans data_tableau, en-tête compris) de la ligne "Bénéfice
        # net" - le chiffre le plus important du rapport, mis en évidence.
        index_benefice_net = 5

        data_tableau = [["Indicateur", "Valeur"]] + [[libelle, valeur] for libelle, valeur in lignes]
        tableau = Table(data_tableau, colWidths=[largeur_disponible * 0.6, largeur_disponible * 0.4])

        style_tableau = [
            ('BACKGROUND', (0, 0), (-1, 0), PDF_BLEU_PRINCIPAL),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, PDF_GRIS_BORDURE),
            # Lignes alternées (à partir de la première ligne de données,
            # l'en-tête - déjà en bleu - n'entre pas dans l'alternance).
            *[
                ('BACKGROUND', (0, i), (-1, i), PDF_GRIS_LIGNE_ALTERNEE)
                for i in range(1, len(data_tableau)) if i % 2 == 0
            ],
            ('FONTNAME', (0, index_benefice_net), (-1, index_benefice_net), 'Helvetica-Bold'),
            (
                'TEXTCOLOR', (0, index_benefice_net), (-1, index_benefice_net),
                PDF_VERT_POSITIF if resultat['benefice_net'] >= 0 else PDF_ROUGE_NEGATIF,
            ),
        ]
        tableau.setStyle(TableStyle(style_tableau))
        return tableau

    def _pied_de_page(self, canvas_pdf, doc):
        canvas_pdf.saveState()
        canvas_pdf.setFont('Helvetica', 8)
        canvas_pdf.setFillColor(PDF_GRIS_TEXTE)
        texte = f"Généré par iwiShop le {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
        canvas_pdf.drawString(2 * cm, 1.2 * cm, texte)
        canvas_pdf.restoreState()


class ResumeFinancierExportExcelView(BoutiqueScopedMixin, APIView):
    # Même règle que l'export PDF (ResumeFinancierExportPDFView) : réservé
    # au propriétaire - une extraction de données comptables complètes,
    # pas un simple affichage (P1 point 6, RBAC).
    permission_classes = [IsAuthenticated, IsOwner]

    def get(self, request, *args, **kwargs):
        boutique = self._boutique_effective()
        date_debut = request.GET.get('date_debut')
        date_fin = request.GET.get('date_fin')
        resultat = calculer_resume_financier(boutique, date_debut, date_fin)
        devise = parametres_boutique(boutique).devise

        classeur = Workbook()
        feuille = classeur.active
        feuille.title = "Résumé financier"

        gras = Font(bold=True)

        feuille["A1"] = "Boutique"
        feuille["B1"] = boutique.nom
        feuille["A1"].font = gras
        if resultat['date_debut'] and resultat['date_fin']:
            feuille["A2"] = "Période"
            feuille["B2"] = f"du {resultat['date_debut']} au {resultat['date_fin']}"
        else:
            feuille["A2"] = "Période"
            feuille["B2"] = "toutes dates confondues"
        feuille["A2"].font = gras

        ligne_entete_tableau = 4
        feuille.cell(row=ligne_entete_tableau, column=1, value="Indicateur").font = gras
        feuille.cell(row=ligne_entete_tableau, column=2, value="Valeur").font = gras

        lignes = [
            ("Chiffre d'affaires", f"{resultat['chiffre_affaires']:.2f} {devise}"),
            ("Total des achats", f"{resultat['total_achats']:.2f} {devise}"),
            ("Total des dépenses", f"{resultat['total_depenses']:.2f} {devise}"),
            ("Bénéfice brut", f"{resultat['benefice_brut']:.2f} {devise}"),
            ("Bénéfice net", f"{resultat['benefice_net']:.2f} {devise}"),
            ("Nombre de ventes", resultat['nombre_ventes']),
            ("Nombre d'achats", resultat['nombre_achats']),
            ("Nombre de dépenses", resultat['nombre_depenses']),
        ]
        for decalage, (libelle, valeur) in enumerate(lignes, start=1):
            feuille.cell(row=ligne_entete_tableau + decalage, column=1, value=libelle)
            feuille.cell(row=ligne_entete_tableau + decalage, column=2, value=valeur)

        # Largeur de colonnes ajustée au contenu le plus long de chaque
        # colonne - fichier de travail, pas besoin du même soin visuel que
        # le PDF (pas de couleurs/alternance de lignes).
        for indice_colonne in (1, 2):
            lettre = get_column_letter(indice_colonne)
            plus_long = max(
                len(str(cellule.value)) for cellule in feuille[lettre] if cellule.value is not None
            )
            feuille.column_dimensions[lettre].width = plus_long + 4

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        nom_fichier = f"resume-financier-{boutique.slug}-{timezone.localdate().isoformat()}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'
        classeur.save(response)
        return response
