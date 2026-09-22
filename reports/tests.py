from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.reverse import reverse
from rest_framework.test import APITestCase

from tenants.models import Abonnement, Boutique, FormuleAbonnement, Profil
from suppliers.models import Fournisseur
from products.models import Produit, UniteVente, ProduitPrix
from sales.models import Vente, LigneVente
from parametres.models import ParametresBoutique
from .views import calculer_resume_financier


class ResumeFinancierAccesTests(APITestCase):
    """P2 point 14 : ResumeFinancierView doit réutiliser la résolution de
    boutique effective et le contrôle d'accès de BoutiqueScopedMixin
    (Vue Support comprise) au lieu d'une logique ad-hoc dupliquée."""

    def setUp(self):
        self.boutique = Boutique.objects.create(nom="Boutique A", slug="boutique-a")
        self.user = User.objects.create_user(username="user_a", password="pass1234")
        Profil.objects.create(user=self.user, boutique=self.boutique, est_proprietaire=True)
        self.url = '/api/reports/resume-financier/'

    def test_acces_normal_autorise(self):
        self.client.force_authenticate(user=self.user)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_boutique_desactivee_autorisee(self):
        """Vue en lecture seule : consulter ses rapports reste possible
        boutique désactivée - seules les écritures sont bloquées (scinde
        le contrôle lecture/écriture, audit complémentaire point 1)."""
        self.boutique.actif = False
        self.boutique.save()
        self.client.force_authenticate(user=self.user)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_abonnement_expire_autorise(self):
        """Vue en lecture seule : consulter ses rapports reste possible
        abonnement expiré - seules les écritures sont bloquées (scinde
        le contrôle lecture/écriture, audit complémentaire point 1)."""
        formule = FormuleAbonnement.objects.create(nom="Standard", duree_jours=30, prix=5000)
        Abonnement.objects.create(
            boutique=self.boutique, formule=formule,
            date_debut=timezone.localdate() - timezone.timedelta(days=40),
            date_fin=timezone.localdate() - timezone.timedelta(days=10),
            statut='EXPIRE',
        )
        self.client.force_authenticate(user=self.user)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_vue_support_superuser_autorisee(self):
        admin = User.objects.create_superuser(username="admin", email="admin@example.com", password="x")
        self.client.force_authenticate(user=admin)
        response = self.client.get(self.url, HTTP_X_SUPPORT_BOUTIQUE=str(self.boutique.id))
        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_vue_support_boutique_introuvable_403_pas_500(self):
        """Avant la correction, un ID de Vue Support inexistant faisait
        planter la vue (Boutique.DoesNotExist non gérée -> 500)."""
        admin = User.objects.create_superuser(username="admin2", email="admin2@example.com", password="x")
        self.client.force_authenticate(user=admin)
        response = self.client.get(self.url, HTTP_X_SUPPORT_BOUTIQUE='999999')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_vue_support_ignoree_pour_non_superuser(self):
        """Un employé ne peut pas utiliser l'en-tête Vue Support pour
        consulter le résumé financier d'une autre boutique."""
        autre_boutique = Boutique.objects.create(nom="Boutique B", slug="boutique-b")
        self.client.force_authenticate(user=self.user)
        response = self.client.get(self.url, HTTP_X_SUPPORT_BOUTIQUE=str(autre_boutique.id))
        # L'en-tête est ignoré (non-superuser) : retombe sur sa propre boutique.
        self.assertEqual(response.status_code, status.HTTP_200_OK)


class ResumeFinancierBeneficeNegatifTests(APITestCase):
    """Audit point 3 : les nouveaux validateurs >= 0 sur les prix
    n'empêchent pas un bénéfice CALCULÉ négatif, qui reste légitime quand
    prix_achat > prix de vente (erreur de tarification, vente à perte
    assumée...). Seule la SAISIE d'un prix négatif est bloquée, pas le
    résultat d'un calcul (voir products.tests.ProduitPrixNegatifTests)."""

    def setUp(self):
        self.boutique = Boutique.objects.create(nom="Boutique", slug="boutique-benefice-negatif")
        self.user = User.objects.create_user(username="user", password="pass1234")
        Profil.objects.create(user=self.user, boutique=self.boutique, est_proprietaire=True)
        self.client.force_authenticate(user=self.user)

        self.unite = UniteVente.objects.create(
            boutique=self.boutique, nom="Unité", facteur_conversion=Decimal("1.000"), est_systeme=True
        )
        # Prix d'achat volontairement supérieur au prix de vente.
        self.produit = Produit.objects.create(
            boutique=self.boutique, nom="Produit vendu à perte",
            prix_achat=Decimal("100.00"), prix_unitaire=Decimal("60.00"), prix_douzaine=Decimal("700.00"),
            quantite_en_stock=10,
        )
        ProduitPrix.objects.create(produit=self.produit, unite=self.unite, prix=Decimal("60.00"))

        vente = Vente.objects.create(
            boutique=self.boutique, montant_paye=Decimal("60.00"),
            montant_total=Decimal("60.00"), montant_net=Decimal("60.00"),
        )
        LigneVente.objects.create(
            boutique=self.boutique, vente=vente, produit=self.produit, quantite=1,
            type_vente='UNITE', unite=self.unite, facteur_conversion_applique=Decimal("1.000"),
            prix_applique=Decimal("60.00"),
        )

    def test_benefice_calcule_negatif_reste_autorise(self):
        response = self.client.get('/api/reports/resume-financier/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        benefice_brut = Decimal(str(response.data['benefice_brut']))
        self.assertEqual(benefice_brut, Decimal("-40.00"))


class CoutHistoriqueBeneficeTests(APITestCase):
    """BUG P1 : Produit.prix_achat change à chaque nouvel achat (dernier
    prix payé au fournisseur - voir purchases.serializers.AchatSerializer),
    et ne doit plus jamais être relu pour recalculer le bénéfice d'une vente
    déjà réalisée. Le bénéfice historique doit utiliser
    LigneVente.prix_achat_unitaire, figé au moment de chaque vente.

    Lien avec la limite déjà connue (audit Point 4 - Achats) : annuler un
    achat restaure le stock mais pas Produit.prix_achat (pas d'historique de
    prix côté achats). Ce chantier ne corrige pas cette limite - il l'évite
    simplement côté ventes en ne dépendant plus jamais de la valeur courante
    de Produit.prix_achat pour un bénéfice déjà calculé. Les deux limites
    restent indépendantes : celle des achats concerne l'ANNULATION d'un
    achat, celle-ci concernait la LECTURE d'un coût par les rapports."""

    def setUp(self):
        self.boutique = Boutique.objects.create(nom="Boutique", slug="boutique-cout-historique-rapports")
        self.user = User.objects.create_user(username="user_cout_rapports", password="pass1234")
        Profil.objects.create(user=self.user, boutique=self.boutique, est_proprietaire=True)
        self.client.force_authenticate(user=self.user)

        self.unite = UniteVente.objects.create(
            boutique=self.boutique, nom="Unité", facteur_conversion=Decimal("1.000"), est_systeme=True
        )
        self.produit = Produit.objects.create(
            boutique=self.boutique, nom="Produit",
            prix_achat=Decimal("500.00"), prix_unitaire=Decimal("800.00"), prix_douzaine=Decimal("9600.00"),
            quantite_en_stock=100,
        )
        ProduitPrix.objects.create(produit=self.produit, unite=self.unite, prix=Decimal("800.00"))
        self.fournisseur = Fournisseur.objects.create(boutique=self.boutique, nom="Fournisseur")

    def _vendre(self, quantite):
        response = self.client.post(reverse('ventes-list'), {
            "montant_paye": str(Decimal("800.00") * quantite),
            "lignes": [{"produit": self.produit.id, "quantite": quantite, "type_vente": "UNITE", "prix_applique": "800.00"}],
        }, format='json')
        assert response.status_code == status.HTTP_201_CREATED, response.data
        return response.data

    def _acheter(self, prix_unitaire_achat):
        response = self.client.post(reverse('achats-list'), {
            "fournisseur": self.fournisseur.id,
            "lignes": [{"produit": self.produit.id, "quantite": 1, "unite": self.unite.id, "prix_unitaire_achat": str(prix_unitaire_achat)}],
        }, format='json')
        assert response.status_code == status.HTTP_201_CREATED, response.data
        return response.data

    def _benefice_brut(self):
        response = self.client.get('/api/reports/resume-financier/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return Decimal(str(response.data['benefice_brut']))

    def test_1_benefice_historique_inchange_apres_nouvel_achat(self):
        self._vendre(10)
        self.assertEqual(self._benefice_brut(), Decimal("3000.00"))  # 10 x (800-500)

        self._acheter(Decimal("700.00"))
        self.produit.refresh_from_db()
        self.assertEqual(self.produit.prix_achat, Decimal("700.00"))

        # Doit rester 3000, pas retomber à 1000 (10 x (800-700)).
        self.assertEqual(self._benefice_brut(), Decimal("3000.00"))

    def test_2_deux_ventes_deux_couts_distincts(self):
        self._vendre(10)                    # coût 500 -> bénéfice 3000
        self._acheter(Decimal("700.00"))
        self._vendre(10)                    # coût 700 -> bénéfice 1000

        lignes = LigneVente.objects.filter(boutique=self.boutique).order_by('id')
        self.assertEqual(lignes.count(), 2)
        self.assertEqual(lignes[0].prix_achat_unitaire, Decimal("500.00"))
        self.assertEqual(lignes[1].prix_achat_unitaire, Decimal("700.00"))
        self.assertEqual(self._benefice_brut(), Decimal("4000.00"))

    def test_3_changement_ulterieur_du_produit_sans_impact_retroactif(self):
        self._vendre(10)
        self._acheter(Decimal("700.00"))
        self._vendre(10)
        self.assertEqual(self._benefice_brut(), Decimal("4000.00"))

        self._acheter(Decimal("900.00"))
        self.produit.refresh_from_db()
        self.assertEqual(self.produit.prix_achat, Decimal("900.00"))
        self.assertEqual(self._benefice_brut(), Decimal("4000.00"))  # toujours 4000

    def test_lignes_pre_migration_sans_cout_historique_retombent_sur_produit(self):
        """Repli documenté (Coalesce) pour les lignes antérieures à ce
        correctif : prix_achat_unitaire NULL -> Produit.prix_achat courant,
        comme avant. Comportement inchangé pour les données existantes,
        jamais utilisé pour une ligne créée après ce correctif."""
        vente = Vente.objects.create(
            boutique=self.boutique, montant_paye=Decimal("800.00"),
            montant_total=Decimal("800.00"), montant_net=Decimal("800.00"),
        )
        LigneVente.objects.create(
            boutique=self.boutique, vente=vente, produit=self.produit, quantite=1,
            type_vente='UNITE', unite=self.unite, facteur_conversion_applique=Decimal("1.000"),
            prix_applique=Decimal("800.00"),
            # prix_achat_unitaire omis : simule une ligne créée avant le correctif.
        )
        self.assertEqual(self._benefice_brut(), Decimal("300.00"))  # 800 - 500 (prix_achat courant)


class ResumeFinancierRemiseBeneficeTests(APITestCase):
    def setUp(self):
        self.boutique = Boutique.objects.create(nom="Boutique remises", slug="boutique-remises-rapports")
        self.user = User.objects.create_user(username="user_remises_rapports", password="pass1234")
        Profil.objects.create(user=self.user, boutique=self.boutique, est_proprietaire=True)
        self.client.force_authenticate(user=self.user)

        self.unite = UniteVente.objects.create(
            boutique=self.boutique, nom="Unité", facteur_conversion=Decimal("1.000"), est_systeme=True
        )
        self.produit = Produit.objects.create(
            boutique=self.boutique, nom="Produit remisé",
            prix_achat=Decimal("900.00"), prix_unitaire=Decimal("1000.00"), prix_douzaine=Decimal("12000.00"),
            quantite_en_stock=100,
        )

    def _creer_vente(self, remise=Decimal("0.00"), statut="VALIDEE"):
        montant_total = Decimal("10000.00")
        montant_net = montant_total - remise
        vente = Vente.objects.create(
            boutique=self.boutique,
            montant_paye=montant_net,
            montant_total=montant_total,
            remise=remise,
            montant_net=montant_net,
            statut=statut,
        )
        LigneVente.objects.create(
            boutique=self.boutique, vente=vente, produit=self.produit, quantite=10,
            type_vente="UNITE", unite=self.unite, facteur_conversion_applique=Decimal("1.000"),
            prix_applique=Decimal("1000.00"), prix_achat_unitaire=Decimal("600.00"),
        )

    def _resume(self):
        response = self.client.get('/api/reports/resume-financier/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response.data

    def test_vente_sans_remise_utilise_le_ca_net_et_le_cout_historique(self):
        self._creer_vente()

        resume = self._resume()
        self.assertEqual(Decimal(str(resume['chiffre_affaires'])), Decimal("10000.00"))
        self.assertEqual(Decimal(str(resume['benefice_brut'])), Decimal("4000.00"))

    def test_vente_avec_remise_diminue_le_benefice_brut(self):
        self._creer_vente(remise=Decimal("1000.00"))

        resume = self._resume()
        self.assertEqual(Decimal(str(resume['chiffre_affaires'])), Decimal("9000.00"))
        self.assertEqual(Decimal(str(resume['benefice_brut'])), Decimal("3000.00"))

    def test_plusieurs_ventes_avec_et_sans_remise_sont_totalisees(self):
        self._creer_vente()
        self._creer_vente(remise=Decimal("1000.00"))

        resume = self._resume()
        self.assertEqual(Decimal(str(resume['chiffre_affaires'])), Decimal("19000.00"))
        self.assertEqual(Decimal(str(resume['benefice_brut'])), Decimal("7000.00"))

    def test_vente_annulee_ne_compte_ni_dans_le_ca_ni_dans_le_benefice(self):
        self._creer_vente(remise=Decimal("1000.00"), statut="ANNULEE")

        resume = self._resume()
        self.assertEqual(Decimal(str(resume['chiffre_affaires'])), Decimal("0.00"))
        self.assertEqual(Decimal(str(resume['benefice_brut'])), Decimal("0.00"))


class CalculerResumeFinancierFonctionTests(APITestCase):
    """Non-régression de l'extraction de la logique de calcul hors de
    ResumeFinancierView.get() (V2 étape 15, export PDF) :
    calculer_resume_financier() doit retourner exactement les mêmes
    valeurs que la vue JSON, réutilisée telle quelle par celle-ci."""

    def setUp(self):
        self.boutique = Boutique.objects.create(nom="Boutique Fonction", slug="boutique-fonction-resume")
        self.user = User.objects.create_user(username="user_fonction_resume", password="pass1234")
        Profil.objects.create(user=self.user, boutique=self.boutique, est_proprietaire=True)
        self.client.force_authenticate(user=self.user)

        self.unite = UniteVente.objects.create(
            boutique=self.boutique, nom="Unité", facteur_conversion=Decimal("1.000"), est_systeme=True
        )
        self.produit = Produit.objects.create(
            boutique=self.boutique, nom="Produit",
            prix_achat=Decimal("500.00"), prix_unitaire=Decimal("800.00"), prix_douzaine=Decimal("9600.00"),
            quantite_en_stock=10,
        )
        ProduitPrix.objects.create(produit=self.produit, unite=self.unite, prix=Decimal("800.00"))

        vente = Vente.objects.create(
            boutique=self.boutique, montant_paye=Decimal("800.00"),
            montant_total=Decimal("800.00"), montant_net=Decimal("800.00"),
        )
        LigneVente.objects.create(
            boutique=self.boutique, vente=vente, produit=self.produit, quantite=1,
            type_vente='UNITE', unite=self.unite, facteur_conversion_applique=Decimal("1.000"),
            prix_applique=Decimal("800.00"), prix_achat_unitaire=Decimal("500.00"),
        )

    def test_fonction_retourne_les_memes_valeurs_que_la_vue_json(self):
        resultat_fonction = calculer_resume_financier(self.boutique, None, None)

        response = self.client.get('/api/reports/resume-financier/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        for cle in (
            'chiffre_affaires', 'total_achats', 'total_depenses', 'benefice_brut',
            'benefice_net', 'nombre_ventes', 'nombre_achats', 'nombre_depenses',
        ):
            self.assertEqual(Decimal(str(resultat_fonction[cle])), Decimal(str(response.data[cle])))


class ResumeFinancierExportPDFTests(APITestCase):
    """V2 étape 15 : export PDF du résumé financier, réservé au propriétaire
    et isolé par boutique (même résolution _boutique_effective que
    ResumeFinancierView - aucun paramètre ne permet de cibler une autre
    boutique que la sienne)."""

    def setUp(self):
        self.boutique_a = Boutique.objects.create(nom="Boutique Export A", slug="boutique-export-a")
        self.proprietaire_a = User.objects.create_user(username="export_proprio_a", password="pass1234")
        Profil.objects.create(user=self.proprietaire_a, boutique=self.boutique_a, est_proprietaire=True)

        self.employe_a = User.objects.create_user(username="export_employe_a", password="pass1234")
        Profil.objects.create(user=self.employe_a, boutique=self.boutique_a, est_proprietaire=False)

        self.boutique_b = Boutique.objects.create(nom="Boutique Export B", slug="boutique-export-b")
        self.proprietaire_b = User.objects.create_user(username="export_proprio_b", password="pass1234")
        Profil.objects.create(user=self.proprietaire_b, boutique=self.boutique_b, est_proprietaire=True)

        self.url = '/api/reports/resume-financier/export-pdf/'

    def test_export_pdf_autorise_pour_le_proprietaire(self):
        self.client.force_authenticate(user=self.proprietaire_a)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('boutique-export-a', response['Content-Disposition'])

    def test_export_pdf_refuse_pour_un_employe(self):
        self.client.force_authenticate(user=self.employe_a)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_export_pdf_isole_par_boutique(self):
        """Un propriétaire ne peut exporter que le résumé de SA propre
        boutique : aucun paramètre de la requête ne permet de cibler une
        autre boutique (résolution via _boutique_effective, pas un id
        passé par le client)."""
        self.client.force_authenticate(user=self.proprietaire_b)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('boutique-export-b', response['Content-Disposition'])
        self.assertNotIn('boutique-export-a', response['Content-Disposition'])

    def test_export_pdf_utilise_la_devise_de_la_boutique_pas_fcfa_en_dur(self):
        ParametresBoutique.objects.create(boutique=self.boutique_a, devise="EUR")

        self.client.force_authenticate(user=self.proprietaire_a)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(b'EUR', response.content)
        self.assertNotIn(b'FCFA', response.content)


class ResumeFinancierExportExcelTests(APITestCase):
    """V2 étape 16 : export Excel du résumé financier, en miroir de
    l'export PDF (ResumeFinancierExportPDFTests) - mêmes règles de
    permission et d'isolation. Le contenu est vérifié via l'API openpyxl
    (format binaire structuré), pas par une recherche de texte brut comme
    pour le PDF."""

    def setUp(self):
        self.boutique_a = Boutique.objects.create(nom="Boutique Excel A", slug="boutique-excel-a")
        self.proprietaire_a = User.objects.create_user(username="excel_proprio_a", password="pass1234")
        Profil.objects.create(user=self.proprietaire_a, boutique=self.boutique_a, est_proprietaire=True)

        self.employe_a = User.objects.create_user(username="excel_employe_a", password="pass1234")
        Profil.objects.create(user=self.employe_a, boutique=self.boutique_a, est_proprietaire=False)

        self.boutique_b = Boutique.objects.create(nom="Boutique Excel B", slug="boutique-excel-b")
        self.proprietaire_b = User.objects.create_user(username="excel_proprio_b", password="pass1234")
        Profil.objects.create(user=self.proprietaire_b, boutique=self.boutique_b, est_proprietaire=True)

        self.url = '/api/reports/resume-financier/export-excel/'

    def _classeur(self, response):
        return load_workbook(BytesIO(response.content))

    def test_export_excel_autorise_pour_le_proprietaire(self):
        self.client.force_authenticate(user=self.proprietaire_a)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('boutique-excel-a', response['Content-Disposition'])

    def test_export_excel_refuse_pour_un_employe(self):
        self.client.force_authenticate(user=self.employe_a)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_export_excel_isole_par_boutique(self):
        """Un propriétaire ne peut exporter que le résumé de SA propre
        boutique, même principe que ResumeFinancierExportPDFTests
        (résolution via _boutique_effective, pas un id passé par le
        client)."""
        self.client.force_authenticate(user=self.proprietaire_b)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('boutique-excel-b', response['Content-Disposition'])
        self.assertNotIn('boutique-excel-a', response['Content-Disposition'])

    def test_export_excel_contient_les_bonnes_valeurs(self):
        ParametresBoutique.objects.create(boutique=self.boutique_a, devise="EUR")
        unite = UniteVente.objects.create(
            boutique=self.boutique_a, nom="Unité", facteur_conversion=Decimal("1.000"), est_systeme=True
        )
        produit = Produit.objects.create(
            boutique=self.boutique_a, nom="Produit Excel",
            prix_achat=Decimal("500.00"), prix_unitaire=Decimal("800.00"), prix_douzaine=Decimal("9600.00"),
            quantite_en_stock=10,
        )
        ProduitPrix.objects.create(produit=produit, unite=unite, prix=Decimal("800.00"))
        vente = Vente.objects.create(
            boutique=self.boutique_a, montant_paye=Decimal("800.00"),
            montant_total=Decimal("800.00"), montant_net=Decimal("800.00"),
        )
        LigneVente.objects.create(
            boutique=self.boutique_a, vente=vente, produit=produit, quantite=1,
            type_vente='UNITE', unite=unite, facteur_conversion_applique=Decimal("1.000"),
            prix_applique=Decimal("800.00"), prix_achat_unitaire=Decimal("500.00"),
        )

        self.client.force_authenticate(user=self.proprietaire_a)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        classeur = self._classeur(response)
        feuille = classeur["Résumé financier"]

        self.assertEqual(feuille["A1"].value, "Résumé financier — Boutique Excel A")

        # Lignes du tableau (après titre fusionné ligne 1, période ligne 2,
        # une ligne 3 vide, en-tête Indicateur/Valeur ligne 4) : lignes 5-12.
        valeurs = {
            feuille.cell(row=r, column=1).value: feuille.cell(row=r, column=2).value
            for r in range(5, 13)
        }
        # Montants écrits comme de vrais nombres (recalculables dans
        # Excel), pas du texte formaté - la devise vit dans number_format.
        self.assertEqual(valeurs["Chiffre d'affaires"], 800.0)
        self.assertEqual(valeurs["Bénéfice brut"], 300.0)
        self.assertEqual(valeurs["Bénéfice net"], 300.0)
        self.assertEqual(valeurs["Nombre de ventes"], 1)

        cellule_ca = feuille.cell(row=5, column=2)
        self.assertIn('EUR', cellule_ca.number_format)
        self.assertIn('#,##0.00', cellule_ca.number_format)

        # Bénéfice net (ligne 9) mis en évidence, même logique que le PDF.
        cellule_benefice_net = feuille.cell(row=9, column=2)
        self.assertTrue(cellule_benefice_net.font.bold)


class VentesDetailleesTestsBase(APITestCase):
    """Fixtures communes aux tests du rapport détaillé des ventes."""

    def setUp(self):
        self.boutique = Boutique.objects.create(nom="Boutique Ventes Détail", slug="boutique-ventes-detail")
        self.proprietaire = User.objects.create_user(username="ventesdetail_proprio", password="pass1234")
        Profil.objects.create(user=self.proprietaire, boutique=self.boutique, est_proprietaire=True)
        self.employe = User.objects.create_user(username="ventesdetail_employe", password="pass1234")
        Profil.objects.create(user=self.employe, boutique=self.boutique, est_proprietaire=False)

        self.unite = UniteVente.objects.create(
            boutique=self.boutique, nom="Unité", facteur_conversion=Decimal("1.000"), est_systeme=True
        )
        self.produit = Produit.objects.create(
            boutique=self.boutique, nom="Produit Détail",
            prix_achat=Decimal("500.00"), prix_unitaire=Decimal("800.00"), prix_douzaine=Decimal("9600.00"),
            quantite_en_stock=100,
        )
        ProduitPrix.objects.create(produit=self.produit, unite=self.unite, prix=Decimal("800.00"))

    def _creer_vente_avec_ligne(self, date_vente, quantite=1, utilisateur=None, statut='VALIDEE'):
        montant = Decimal("800.00") * quantite
        vente = Vente.objects.create(
            boutique=self.boutique, montant_paye=montant, montant_total=montant, montant_net=montant,
            utilisateur=utilisateur, statut=statut,
        )
        LigneVente.objects.create(
            boutique=self.boutique, vente=vente, produit=self.produit, quantite=quantite,
            type_vente='UNITE', unite=self.unite, facteur_conversion_applique=Decimal("1.000"),
            prix_applique=Decimal("800.00"), prix_achat_unitaire=Decimal("500.00"),
        )
        Vente.objects.filter(pk=vente.pk).update(date_vente=date_vente)
        vente.refresh_from_db()
        return vente

    def _creer_vente_pour_boutique(self, boutique, nom_produit, date_vente):
        """Comme _creer_vente_avec_ligne, mais pour une boutique arbitraire
        (pas forcément self.boutique) - utilisé par les tests d'isolation
        multi-tenant qui doivent comparer le contenu d'un export entre deux
        boutiques distinctes."""
        # get_or_create (pas create) : quand boutique=self.boutique, une
        # UniteVente "Unité" existe déjà (créée par
        # VentesDetailleesTestsBase.setUp) - la recréer violerait la
        # contrainte unique (boutique, nom).
        unite, _ = UniteVente.objects.get_or_create(
            boutique=boutique, nom="Unité",
            defaults={"facteur_conversion": Decimal("1.000"), "est_systeme": True},
        )
        produit = Produit.objects.create(
            boutique=boutique, nom=nom_produit,
            prix_achat=Decimal("500.00"), prix_unitaire=Decimal("800.00"), prix_douzaine=Decimal("9600.00"),
            quantite_en_stock=100,
        )
        vente = Vente.objects.create(
            boutique=boutique, montant_paye=Decimal("800.00"), montant_total=Decimal("800.00"),
            montant_net=Decimal("800.00"), statut='VALIDEE',
        )
        LigneVente.objects.create(
            boutique=boutique, vente=vente, produit=produit, quantite=1,
            type_vente='UNITE', unite=unite, facteur_conversion_applique=Decimal("1.000"),
            prix_applique=Decimal("800.00"), prix_achat_unitaire=Decimal("500.00"),
        )
        Vente.objects.filter(pk=vente.pk).update(date_vente=date_vente)
        return vente


class VentesDetailleesViewTests(VentesDetailleesTestsBase):
    def setUp(self):
        super().setUp()
        self.url = reverse('ventes-detaillees')

    def test_bornes_dates_obligatoires(self):
        self.client.force_authenticate(user=self.proprietaire)

        self.assertEqual(self.client.get(self.url).status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            self.client.get(self.url, {"date_debut": "2026-09-01"}).status_code, status.HTTP_400_BAD_REQUEST
        )

    def test_employe_peut_lire(self):
        """Même principe que ResumeFinancierView : lecture ouverte à
        l'employé, seuls les exports sont réservés au propriétaire."""
        self.client.force_authenticate(user=self.employe)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)

    def test_ne_retourne_que_les_lignes_dans_la_periode_et_validees(self):
        vente_dans_periode = self._creer_vente_avec_ligne(
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)), utilisateur=self.proprietaire,
        )
        self._creer_vente_avec_ligne(timezone.make_aware(timezone.datetime(2026, 8, 1, 10, 0)))
        vente_annulee = self._creer_vente_avec_ligne(
            timezone.make_aware(timezone.datetime(2026, 9, 20, 10, 0)), statut='ANNULEE',
        )

        self.client.force_authenticate(user=self.proprietaire)
        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        numeros = [ligne['numero_vente'] for ligne in response.data['lignes']]
        self.assertIn(vente_dans_periode.numero, numeros)
        self.assertNotIn(vente_annulee.numero, numeros)
        self.assertEqual(len(numeros), 1)

    def test_contenu_dune_ligne(self):
        vente = self._creer_vente_avec_ligne(
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)), quantite=3, utilisateur=self.proprietaire,
        )

        self.client.force_authenticate(user=self.proprietaire)
        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        ligne = response.data['lignes'][0]
        self.assertEqual(ligne['numero_vente'], vente.numero)
        self.assertEqual(ligne['vendeur'], self.proprietaire.username)
        self.assertEqual(ligne['produit'], "Produit Détail")
        self.assertEqual(ligne['quantite'], 3)
        self.assertEqual(ligne['unite'], "Unité")
        self.assertEqual(Decimal(str(ligne['prix_applique'])), Decimal("800.00"))
        self.assertEqual(Decimal(str(ligne['sous_total'])), Decimal("2400.00"))


class VentesDetailleesRequeteCountTests(VentesDetailleesTestsBase):
    """select_related doit éviter tout N+1 : le nombre de requêtes ne doit
    pas dépendre du nombre de lignes retournées (même pattern que
    sales.tests.VenteListQueryCountTests)."""

    def setUp(self):
        super().setUp()
        self.url = reverse('ventes-detaillees')
        self.client.force_authenticate(user=self.proprietaire)

    def test_nombre_de_requetes_constant_quel_que_soit_le_nombre_de_lignes(self):
        self._creer_vente_avec_ligne(
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)), utilisateur=self.proprietaire,
        )
        with CaptureQueriesContext(connection) as premier:
            response_1 = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        for _ in range(5):
            self._creer_vente_avec_ligne(
                timezone.make_aware(timezone.datetime(2026, 9, 16, 10, 0)), utilisateur=self.proprietaire,
            )
        with CaptureQueriesContext(connection) as second:
            response_2 = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response_1.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response_2.data['lignes']), 6)
        self.assertEqual(len(premier.captured_queries), len(second.captured_queries))


class VentesDetailleesExportPDFTests(VentesDetailleesTestsBase):
    def setUp(self):
        super().setUp()
        self.autre_boutique = Boutique.objects.create(nom="Autre Boutique Détail", slug="autre-boutique-detail")
        self.autre_proprietaire = User.objects.create_user(username="ventesdetail_autre_proprio", password="pass1234")
        Profil.objects.create(user=self.autre_proprietaire, boutique=self.autre_boutique, est_proprietaire=True)
        self.url = reverse('ventes-detaillees-export-pdf')

    def test_bornes_dates_obligatoires(self):
        self.client.force_authenticate(user=self.proprietaire)
        self.assertEqual(self.client.get(self.url).status_code, status.HTTP_400_BAD_REQUEST)

    def test_autorise_pour_le_proprietaire(self):
        self.client.force_authenticate(user=self.proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('boutique-ventes-detail', response['Content-Disposition'])

    def test_refuse_pour_un_employe(self):
        self.client.force_authenticate(user=self.employe)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_isole_par_boutique(self):
        self.client.force_authenticate(user=self.autre_proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('autre-boutique-detail', response['Content-Disposition'])
        self.assertNotIn('boutique-ventes-detail', response['Content-Disposition'])

    def test_contient_les_donnees_de_la_vente(self):
        self._creer_vente_avec_ligne(
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)), utilisateur=self.proprietaire,
        )
        self.client.force_authenticate(user=self.proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(b'Produit', response.content)
        self.assertIn(self.proprietaire.username.encode(), response.content)

    def test_contenu_isole_par_boutique(self):
        """Preuve directe (pas seulement le nom de fichier) : le PDF de
        autre_proprietaire contient le produit de SA boutique mais jamais
        celui de boutique-ventes-detail, même période, même requête."""
        self._creer_vente_pour_boutique(
            self.boutique, "Produit Boutique A Isolation",
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)),
        )
        self._creer_vente_pour_boutique(
            self.autre_boutique, "Produit Boutique B Isolation",
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)),
        )
        self.client.force_authenticate(user=self.autre_proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(b'Produit Boutique B Isolation', response.content)
        self.assertNotIn(b'Produit Boutique A Isolation', response.content)


class VentesDetailleesExportExcelTests(VentesDetailleesTestsBase):
    def setUp(self):
        super().setUp()
        self.autre_boutique = Boutique.objects.create(
            nom="Autre Boutique Excel Détail", slug="autre-boutique-excel-detail"
        )
        self.autre_proprietaire = User.objects.create_user(username="ventesdetail_excel_autre", password="pass1234")
        Profil.objects.create(user=self.autre_proprietaire, boutique=self.autre_boutique, est_proprietaire=True)
        self.url = reverse('ventes-detaillees-export-excel')

    def _classeur(self, response):
        return load_workbook(BytesIO(response.content))

    def test_bornes_dates_obligatoires(self):
        self.client.force_authenticate(user=self.proprietaire)
        self.assertEqual(self.client.get(self.url).status_code, status.HTTP_400_BAD_REQUEST)

    def test_autorise_pour_le_proprietaire(self):
        self.client.force_authenticate(user=self.proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('boutique-ventes-detail', response['Content-Disposition'])

    def test_refuse_pour_un_employe(self):
        self.client.force_authenticate(user=self.employe)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_isole_par_boutique(self):
        self.client.force_authenticate(user=self.autre_proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('autre-boutique-excel-detail', response['Content-Disposition'])

    def test_contenu_dune_ligne(self):
        ParametresBoutique.objects.create(boutique=self.boutique, devise="EUR")
        self._creer_vente_avec_ligne(
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)), quantite=2, utilisateur=self.proprietaire,
        )
        self.client.force_authenticate(user=self.proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        classeur = self._classeur(response)
        feuille = classeur["Ventes détaillées"]

        self.assertEqual(feuille["A1"].value, "Ventes détaillées — Boutique Ventes Détail")
        # En-tête du tableau (ligne 4) puis la seule ligne de données (ligne 5).
        entetes = [feuille.cell(row=4, column=c).value for c in range(1, 10)]
        self.assertEqual(
            entetes,
            ["Date", "N° vente", "Vendeur", "Client", "Produit", "Qté", "Unité", "Prix", "Sous-total"],
        )

        valeurs = [feuille.cell(row=5, column=c).value for c in range(1, 10)]
        self.assertEqual(valeurs[2], self.proprietaire.username)  # Vendeur
        self.assertEqual(valeurs[4], "Produit Détail")            # Produit
        self.assertEqual(valeurs[5], 2)                           # Qté
        self.assertEqual(valeurs[7], 800.0)                       # Prix
        self.assertEqual(valeurs[8], 1600.0)                      # Sous-total

        cellule_sous_total = feuille.cell(row=5, column=9)
        self.assertIn('EUR', cellule_sous_total.number_format)

    def test_contenu_isole_par_boutique(self):
        """Preuve directe (pas seulement le nom de fichier) : le classeur de
        autre_proprietaire contient le produit de SA boutique mais jamais
        celui de boutique-ventes-detail, même période, même requête."""
        self._creer_vente_pour_boutique(
            self.boutique, "Produit Boutique A Isolation",
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)),
        )
        self._creer_vente_pour_boutique(
            self.autre_boutique, "Produit Boutique B Isolation",
            timezone.make_aware(timezone.datetime(2026, 9, 15, 10, 0)),
        )
        self.client.force_authenticate(user=self.autre_proprietaire)

        response = self.client.get(self.url, {"date_debut": "2026-09-01", "date_fin": "2026-09-30"})

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        classeur = self._classeur(response)
        feuille = classeur["Ventes détaillées"]
        produits_lus = [feuille.cell(row=5, column=5).value]

        self.assertIn("Produit Boutique B Isolation", produits_lus)
        self.assertNotIn("Produit Boutique A Isolation", produits_lus)
